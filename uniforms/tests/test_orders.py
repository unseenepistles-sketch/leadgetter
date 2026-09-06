"""The order lifecycle as it actually runs.

She raises an order against a PR number, the tailor visits to take measurements,
and the tailor comes back with the clothes — sometimes all of them, sometimes
some now and the rest later. The renewal clock starts when the garment reaches
the person, not when the order was raised.
"""
from __future__ import annotations

from datetime import date, timedelta

import pytest

from app.domain.models import OrderStatus, order_status
from app.excelstore.workbook import WorkbookStore
from app.service import NotFound, UniformService, ValidationError
from tests.fixtures import clean_workbook


@pytest.fixture
def service(tmp_path):
    store = WorkbookStore(clean_workbook(tmp_path / "w.xlsx"), backup_dir=tmp_path / "bak")
    store.load()
    return UniformService(store)


def lines(*specs):
    """Terser than a dict literal per line, and these tests are full of them."""
    return [
        {"employee_number": e, "item_code": c, "quantity": q} for e, c, q in specs
    ]


# --------------------------------------------------------------- the status rule

@pytest.mark.parametrize(
    "qty,delivered,expected",
    [
        (2, 0, OrderStatus.PENDING),
        (2, 1, OrderStatus.PARTIAL),
        (2, 2, OrderStatus.DELIVERED),
        (2, 3, OrderStatus.DELIVERED),   # over-delivery still counts as done
        (1, 0, OrderStatus.PENDING),
    ],
)
def test_order_status(qty, delivered, expected):
    assert order_status(qty, delivered) is expected


def test_unmeasured_order_is_awaiting_measurement():
    assert order_status(2, 0, measured=False) is OrderStatus.AWAITING_MEASUREMENT


def test_a_delivery_outranks_a_missing_measurement_date():
    """If the clothes are here the tailor plainly came, whatever the sheet says."""
    assert order_status(2, 1, measured=False) is OrderStatus.PARTIAL
    assert order_status(2, 2, measured=False) is OrderStatus.DELIVERED


def test_measurement_defaults_to_recorded():
    """An older workbook with no measurement column behaves exactly as before."""
    assert order_status(2, 0) is OrderStatus.PENDING


def test_cancelled_beats_everything():
    assert order_status(2, 1, cancelled=True) is OrderStatus.CANCELLED
    assert order_status(2, 0, cancelled=True, measured=False) is OrderStatus.CANCELLED


# ------------------------------------------------------------------- raising one

def test_one_pr_covers_many_people_and_many_items(service):
    summary = service.place_order(
        "PR-2026-1077",
        lines(("E002", "SHIRT", 6), ("E002", "TROUSER", 3), ("E003", "SHIRT", 6)),
    )
    assert summary.pr_number == "PR-2026-1077"
    assert len(summary.lines) == 3
    assert summary.people == 2
    assert summary.ordered == 15
    assert {l.order.order_id for l in summary.lines} == {"PR-2026-1077"}


def test_a_new_order_is_awaiting_measurement(service):
    summary = service.place_order("PR-1", lines(("E002", "SHIRT", 2)))
    assert summary.measured is None
    assert summary.status is OrderStatus.AWAITING_MEASUREMENT


def test_the_pr_number_is_required(service):
    with pytest.raises(ValidationError, match="PR number"):
        service.place_order("   ", lines(("E002", "SHIRT", 1)))


def test_a_pr_number_cannot_be_reused(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1)))
    with pytest.raises(ValidationError, match="already been used"):
        service.place_order("PR-1", lines(("E003", "SHIRT", 1)))


def test_the_same_person_and_item_cannot_appear_twice(service):
    with pytest.raises(ValidationError, match="twice"):
        service.place_order("PR-1", lines(("E002", "SHIRT", 2), ("E002", "SHIRT", 4)))


def test_zero_quantities_are_ignored(service):
    summary = service.place_order(
        "PR-1", lines(("E002", "SHIRT", 2), ("E002", "TROUSER", 0)))
    assert [l.order.item_code for l in summary.lines] == ["SHIRT"]


def test_an_order_of_nothing_is_rejected(service):
    with pytest.raises(ValidationError, match="at least one item"):
        service.place_order("PR-1", lines(("E002", "SHIRT", 0)))


def test_unknown_item_is_rejected(service):
    with pytest.raises(ValidationError, match="unknown item"):
        service.place_order("PR-1", lines(("E002", "NOPE", 1)))


def test_unknown_employee_is_rejected(service):
    with pytest.raises(NotFound):
        service.place_order("PR-1", lines(("NOBODY", "SHIRT", 1)))


def test_a_future_dated_order_is_rejected(service):
    with pytest.raises(ValidationError, match="future"):
        service.place_order("PR-1", lines(("E002", "SHIRT", 1)),
                            ordered_date=date.today() + timedelta(days=1))


def test_sizes_ride_along_on_the_order(service):
    summary = service.place_order(
        "PR-1", [{"employee_number": "E002", "item_code": "SHIRT",
                  "quantity": 1, "size": "42"}])
    assert summary.lines[0].order.size == "42"


# ------------------------------------------------------ the measurement visit

def test_recording_the_measurement_moves_it_on(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 2), ("E003", "SHIRT", 2)),
                        ordered_date=date(2026, 1, 5))
    summary = service.record_measurement("PR-1", date(2026, 1, 20))
    assert summary.measured == date(2026, 1, 20)
    assert summary.status is OrderStatus.PENDING


def test_one_visit_covers_every_line_on_the_pr(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1), ("E003", "TROUSER", 1)),
                        ordered_date=date(2026, 1, 5))
    service.record_measurement("PR-1", date(2026, 1, 20))
    assert all(l.order.measured_date == date(2026, 1, 20)
               for l in service.order("PR-1").lines)


def test_measuring_before_the_order_was_raised_is_rejected(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1)), ordered_date=date(2026, 6, 1))
    with pytest.raises(ValidationError, match="before the order was raised"):
        service.record_measurement("PR-1", date(2026, 5, 1))


def test_a_future_measurement_date_is_rejected(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1)))
    with pytest.raises(ValidationError, match="future"):
        service.record_measurement("PR-1", date.today() + timedelta(days=1))


def test_measuring_an_unknown_pr_is_not_found(service):
    with pytest.raises(NotFound):
        service.record_measurement("PR-NOPE")


def test_the_measurement_is_written_to_the_audit_trail(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1)), ordered_date=date(2026, 1, 5))
    before = service.store.pending_writes
    service.record_measurement("PR-1", date(2026, 1, 20), who="njiru")
    assert service.store.pending_writes > before


# ------------------------------------------------------------------ delivery

def measured(service, pr="PR-1", **kw):
    """Raise an order and get the tailor through the door — the usual preamble."""
    summary = service.place_order(pr, kw.pop("order_lines"),
                                  ordered_date=kw.pop("ordered_date", date(2026, 1, 5)))
    service.record_measurement(pr, kw.pop("measured_date", date(2026, 1, 20)))
    return summary


def test_delivery_needs_the_measurement_first(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 2)))
    with pytest.raises(ValidationError, match="measurement visit"):
        service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)))


def test_part_delivery_reports_partially_delivered(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 4)),
                               received_date=date(2026, 2, 10))
    line = service.order_line("PR-1", "E002", "SHIRT")
    assert (line.delivered, line.pending) == (4, 2)
    assert line.status is OrderStatus.PARTIAL


def test_delivering_the_rest_closes_it(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 4)),
                               received_date=date(2026, 2, 10))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 3, 3))
    assert service.order("PR-1").status is OrderStatus.DELIVERED


def test_an_order_is_only_delivered_when_every_line_is(service):
    """One outstanding shirt is the whole reason she is looking at the screen."""
    measured(service, order_lines=lines(("E002", "SHIRT", 2), ("E003", "SHIRT", 2)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 2, 10))
    assert service.order("PR-1").status is OrderStatus.PARTIAL


def test_one_persons_delivery_is_not_credited_to_another(service):
    """The bug the (PR, person, item) key exists to prevent."""
    measured(service, order_lines=lines(("E002", "SHIRT", 2), ("E003", "SHIRT", 2)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 2, 10))
    assert service.order_line("PR-1", "E002", "SHIRT").delivered == 2
    assert service.order_line("PR-1", "E003", "SHIRT").delivered == 0
    assert service.order_line("PR-1", "E003", "SHIRT").pending == 2


def test_delivery_defaults_to_everything_outstanding(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 5)))
    service.receive_delivery("PR-1", "E002", "SHIRT", received_date=date(2026, 2, 10))
    assert service.order_line("PR-1", "E002", "SHIRT").status is OrderStatus.DELIVERED


def test_cannot_deliver_more_than_was_ordered(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    with pytest.raises(ValidationError, match="outstanding"):
        service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 3)))


def test_a_bad_line_writes_none_of_the_delivery(service):
    """Half a delivery recorded is worse than none — validate before writing."""
    measured(service, order_lines=lines(("E002", "SHIRT", 2), ("E003", "SHIRT", 2)))
    with pytest.raises(ValidationError):
        service.receive_deliveries(
            "PR-1", lines(("E002", "SHIRT", 2), ("E003", "SHIRT", 99)))
    assert service.order_line("PR-1", "E002", "SHIRT").delivered == 0


def test_cannot_deliver_a_closed_line(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 2, 10))
    with pytest.raises(ValidationError, match="already fully delivered"):
        service.receive_delivery("PR-1", "E002", "SHIRT", quantity=1)


def test_future_dated_delivery_is_rejected(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    with pytest.raises(ValidationError, match="future"):
        service.receive_delivery("PR-1", "E002", "SHIRT",
                                 received_date=date.today() + timedelta(days=1))


def test_delivery_before_the_order_was_raised_is_rejected(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)),
             ordered_date=date(2026, 6, 1), measured_date=date(2026, 6, 10))
    with pytest.raises(ValidationError, match="before the order was raised"):
        service.receive_delivery("PR-1", "E002", "SHIRT", received_date=date(2026, 5, 1))


def test_delivering_nothing_is_rejected(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    with pytest.raises(ValidationError, match="nothing was marked"):
        service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 0)))


def test_unknown_order_is_not_found(service):
    with pytest.raises(NotFound):
        service.order_line("PR-NOPE", "E002", "SHIRT")
    with pytest.raises(NotFound):
        service.order("PR-NOPE")


# ---------------------------------------------- what delivery does downstream

def test_delivery_records_an_issuance_linked_to_the_pr(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    [issuance] = service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                                            received_date=date(2026, 2, 10))
    assert issuance.order_id == "PR-1"
    assert issuance.issued_date == date(2026, 2, 10)
    assert issuance.employee_number == "E002"


def test_the_ordered_size_carries_onto_the_handover(service):
    service.place_order("PR-1", [{"employee_number": "E002", "item_code": "SHIRT",
                                  "quantity": 1, "size": "42"}],
                        ordered_date=date(2026, 1, 5))
    service.record_measurement("PR-1", date(2026, 1, 20))
    [issuance] = service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 1)),
                                            received_date=date(2026, 2, 10))
    assert issuance.size == "42"


def test_delivery_moves_the_employee_out_of_never_issued(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 1)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 1)),
                               received_date=date(2026, 2, 10))
    statuses = {s.item_code: s.status for s in service.statuses_for("E002")}
    assert statuses["SHIRT"].value != "never_issued"


def test_renewal_clock_runs_from_delivery_not_from_the_order(service):
    """Ordered in January, delivered in June, renews next June."""
    measured(service, order_lines=lines(("E002", "SHIRT", 1)),
             ordered_date=date(2026, 1, 5), measured_date=date(2026, 1, 20))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 1)),
                               received_date=date(2026, 6, 1))
    line = service.order_line("PR-1", "E002", "SHIRT")
    assert line.last_delivery == date(2026, 6, 1)
    assert line.next_due == date(2027, 6, 1)


def test_each_part_delivery_keeps_its_own_date(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 4)),
                               received_date=date(2026, 2, 10))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 3, 3))
    dates = [i.issued_date for i in service.deliveries_for_order("PR-1")]
    assert dates == [date(2026, 2, 10), date(2026, 3, 3)]


def test_pending_delivery_counts_pieces_not_orders(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6), ("E003", "TROUSER", 3)))
    assert service.pending_delivery() == 9


def test_pending_delivery_drops_as_stock_arrives(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 4)),
                               received_date=date(2026, 2, 10))
    assert service.pending_delivery() == 2


# ------------------------------------------------------------------ the log

def test_the_order_log_rolls_up_by_pr(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1), ("E003", "SHIRT", 1)),
                        ordered_date=date(2026, 1, 5))
    service.place_order("PR-2", lines(("E002", "TROUSER", 1)),
                        ordered_date=date(2026, 2, 5))
    log = service.orders()
    assert [o.pr_number for o in log] == ["PR-2", "PR-1"]   # newest first
    assert [o.lines and len(o.lines) for o in log] == [1, 2]


def test_filtering_the_order_log_by_status(service):
    service.place_order("PR-1", lines(("E002", "SHIRT", 1)), ordered_date=date(2026, 1, 5))
    service.place_order("PR-2", lines(("E003", "SHIRT", 1)), ordered_date=date(2026, 2, 5))
    service.record_measurement("PR-2", date(2026, 2, 10))
    assert [o.pr_number for o in service.orders(status="awaiting_measurement")] == ["PR-1"]
    assert [o.pr_number for o in service.orders(status="pending")] == ["PR-2"]


def test_an_unknown_status_filter_is_rejected(service):
    with pytest.raises(ValidationError, match="unknown order status"):
        service.orders(status="banana")


def test_action_needed_describes_part_deliveries(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 6)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 4)),
                               received_date=date(2026, 2, 10))
    rows = service.action_needed(limit=50)
    assert any(r["kind"] == "order" and "4/6" in r["detail"] for r in rows)


def test_action_needed_puts_overdue_above_undelivered(service):
    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    rows = service.action_needed(limit=50)
    ranks = [r["rank"] for r in rows]
    assert ranks == sorted(ranks)


# ------------------------------------------------------------- persistence

def test_orders_survive_a_reload(service, tmp_path):
    service.place_order("PR-2026-1077", lines(("E002", "SHIRT", 6), ("E003", "SHIRT", 6)),
                        ordered_date=date(2026, 1, 5))
    service.record_measurement("PR-2026-1077", date(2026, 1, 20))
    service.store.flush()

    fresh = WorkbookStore(service.store.path, backup_dir=tmp_path / "bak2")
    fresh.load()
    reloaded = UniformService(fresh).order("PR-2026-1077")
    assert reloaded.people == 2
    assert reloaded.ordered == 12
    assert reloaded.measured == date(2026, 1, 20)
    assert reloaded.status is OrderStatus.PENDING


def test_the_pr_number_reaches_the_sheet(service):
    """She types a PR number so she can find it in Excel; it has to be in there."""
    from openpyxl import load_workbook

    service.place_order("PR-2026-1077", lines(("E002", "SHIRT", 1)))
    service.store.flush()
    ws = load_workbook(service.store.path)["Orders"]
    headers = [c.value for c in ws[1]]
    assert "PR Number" in headers
    column = headers.index("PR Number") + 1
    assert any(ws.cell(r, column).value == "PR-2026-1077"
               for r in range(2, ws.max_row + 1))


def test_orders_and_issuances_flush_to_their_own_sheets(service):
    from openpyxl import load_workbook

    measured(service, order_lines=lines(("E002", "SHIRT", 2)))
    service.receive_deliveries("PR-1", lines(("E002", "SHIRT", 2)),
                               received_date=date(2026, 2, 10))
    service.store.flush()
    wb = load_workbook(service.store.path)
    assert wb["Orders"].max_row >= 2
    assert wb["Issuances"].max_row >= 2
