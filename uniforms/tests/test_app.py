"""End-to-end tests through the real ASGI app, workbook and all."""
from __future__ import annotations

import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook

from tests.fixtures import clean_workbook, messy_workbook


@pytest.fixture
def client(tmp_path, monkeypatch):
    workbook = clean_workbook(tmp_path / "uniforms.xlsx")
    monkeypatch.setenv("WORKBOOK_PATH", str(workbook))
    monkeypatch.setenv("BACKUP_DIR", str(tmp_path / "bak"))
    monkeypatch.setenv("LEDGER_PATH", str(tmp_path / "ledger.sqlite3"))
    monkeypatch.setenv("STORES_EMAIL", "stores@co.test")
    monkeypatch.setenv("REMINDERS_ENABLED", "0")

    from app import config, deps

    for cached in (config.get_settings, deps.get_store, deps.get_service, deps.get_ledger):
        cached.cache_clear()

    from app.main import app

    with TestClient(app) as c:
        yield c

    for cached in (config.get_settings, deps.get_store, deps.get_service, deps.get_ledger):
        cached.cache_clear()


# --- pages render ---

@pytest.mark.parametrize(
    "path",
    ["/", "/employees", "/employees/E001", "/issue", "/reports?state=never_issued", "/data-quality"],
)
def test_pages_render(client, path):
    response = client.get(path)
    assert response.status_code == 200
    assert "<table" in response.text or "tile" in response.text


def test_unknown_employee_page_redirects_with_a_message(client):
    response = client.get("/employees/NOPE", follow_redirects=False)
    assert response.status_code == 303
    assert "err=" in response.headers["location"]


# --- the API ---

def test_employee_list_paginates(client):
    body = client.get("/api/employees?per_page=2").json()
    assert body["per_page"] == 2 and body["total"] == 4 and body["pages"] == 2
    assert len(body["items"]) == 2


def test_employee_search(client):
    body = client.get("/api/employees?q=amina").json()
    assert [e["employee_number"] for e in body["items"]] == ["E002"]


def test_active_only_filter_excludes_leavers(client):
    numbers = {e["employee_number"] for e in client.get("/api/employees?active_only=true").json()["items"]}
    assert "E004" not in numbers


def test_never_issued_report_finds_people_with_no_issuance_rows(client):
    """The headline requirement: invisible to any query over the issuance log."""
    body = client.get("/api/status?state=never_issued").json()
    assert body["total"] > 0
    assert any(i["employee_number"] == "E002" for i in body["items"])


def test_status_report_rejects_an_unknown_state(client):
    assert client.get("/api/status?state=banana").status_code == 400


def test_summary_counts_every_status(client):
    counts = client.get("/api/dashboard/summary").json()["counts"]
    assert counts["employees"] == 3  # E004 has left
    assert counts["outstanding"] >= 1


# --- writing through the API ---

def test_record_issuance_prefills_size_and_quantity(client):
    body = client.post(
        "/api/issuances", json={"employee_number": "E002", "item_code": "SHIRT"}
    ).json()
    assert body["size"] == "M"          # from the employee's profile
    assert body["quantity"] == 3        # from the role's entitlement
    assert body["cycle_months"] == 12   # snapshotted at issue time


def test_issuance_moves_the_employee_out_of_never_issued(client):
    before = client.get("/api/employees/E002/uniforms").json()["items"]
    assert {i["status"] for i in before} == {"never_issued"}

    client.post("/api/issuances", json={"employee_number": "E002", "item_code": "SHIRT"})

    after = {i["item_code"]: i["status"] for i in client.get("/api/employees/E002/uniforms").json()["items"]}
    assert after["SHIRT"] == "ok"
    assert after["BLAZER"] == "never_issued"


def test_issuance_reaches_the_workbook_on_disk(client, tmp_path):
    client.post("/api/issuances", json={"employee_number": "E002", "item_code": "SHIRT"})
    assert client.post("/api/workbook/flush").json()["written"] == 1

    wb = load_workbook(tmp_path / "uniforms.xlsx")
    last = list(wb["Issuances"].iter_rows(min_row=2, values_only=True))[-1]
    wb.close()
    assert last[0] == "E002" and last[1] == "SHIRT"


def test_future_dated_issuance_is_rejected(client):
    response = client.post(
        "/api/issuances",
        json={"employee_number": "E002", "item_code": "SHIRT", "issued_date": "2099-01-01"},
    )
    assert response.status_code == 400


def test_unknown_item_is_rejected(client):
    response = client.post(
        "/api/issuances", json={"employee_number": "E002", "item_code": "GHOST"}
    )
    assert response.status_code == 400


def test_unknown_employee_is_a_404(client):
    response = client.post("/api/issuances", json={"employee_number": "NOPE", "item_code": "SHIRT"})
    assert response.status_code == 404


def test_malformed_date_is_rejected(client):
    response = client.post(
        "/api/issuances",
        json={"employee_number": "E002", "item_code": "SHIRT", "issued_date": "15/01/2024"},
    )
    assert response.status_code == 400


def test_bulk_issue_covers_everyone_selected(client):
    body = client.post(
        "/api/issuances/bulk",
        json={"employee_numbers": ["E002", "E003"], "item_codes": ["SHIRT", "TROUSER"]},
    ).json()
    assert body["created"] == 4
    assert {i["employee_number"] for i in body["issuances"]} == {"E002", "E003"}


def test_bulk_issue_needs_both_sides(client):
    assert client.post("/api/issuances/bulk", json={"employee_numbers": [], "item_codes": ["SHIRT"]}).status_code == 400
    assert client.post("/api/issuances/bulk", json={"employee_numbers": ["E002"], "item_codes": []}).status_code == 400


# --- the issue form ---

def test_issue_form_records_and_redirects(client):
    response = client.post(
        "/issue",
        data={"employee_number": "E002", "item_codes": ["SHIRT", "TROUSER"], "issued_by": "stores"},
        follow_redirects=False,
    )
    assert response.status_code == 303
    assert "ok=" in response.headers["location"]
    statuses = {i["item_code"]: i["status"] for i in client.get("/api/employees/E002/uniforms").json()["items"]}
    assert statuses["SHIRT"] == "ok" and statuses["TROUSER"] == "ok"


def test_issue_form_with_no_items_reports_the_error(client):
    response = client.post(
        "/issue", data={"employee_number": "E002"}, follow_redirects=False
    )
    assert "err=" in response.headers["location"]


# --- reminders through the API ---

def test_reminder_dry_run_changes_nothing(client):
    body = client.post("/api/reminders/run?dry_run=true").json()
    assert body["dry_run"] and body["scanned"] > 0
    assert client.get("/api/reminders").json()["counts"] == {}


def test_reminder_live_run_then_repeat_is_silent(client):
    first = client.post("/api/reminders/run?dry_run=false").json()
    assert first["sent"] > 0
    second = client.post("/api/reminders/run?dry_run=false").json()
    assert second["sent"] == 0


# --- data quality surfaces a messy workbook ---

def test_data_quality_reports_problem_rows(tmp_path, monkeypatch):
    workbook = messy_workbook(tmp_path / "messy.xlsx")
    monkeypatch.setenv("WORKBOOK_PATH", str(workbook))
    monkeypatch.setenv("BACKUP_DIR", str(tmp_path / "bak"))
    monkeypatch.setenv("LEDGER_PATH", str(tmp_path / "l.sqlite3"))

    from app import config, deps

    for cached in (config.get_settings, deps.get_store, deps.get_service, deps.get_ledger):
        cached.cache_clear()
    from app.main import app

    with TestClient(app) as client:
        body = client.get("/api/data-quality").json()
        assert body["total"] > 0
        assert client.get("/data-quality").status_code == 200

    for cached in (config.get_settings, deps.get_store, deps.get_service, deps.get_ledger):
        cached.cache_clear()


# --- ordering through the app ---

def test_order_pages_render(client):
    for path in ("/orders", "/orders/new"):
        assert client.get(path).status_code == 200


def test_place_an_order_through_the_form(client):
    response = client.post(
        "/orders",
        data={"pr_number": "PR-2026-1077", "employee_number": ["E002", "E003"],
              "qty_SHIRT": "6", "qty_TROUSER": "3", "qty_BLAZER": "0",
              "tailor": "Al Noor Tailors"},
        follow_redirects=False,
    )
    assert response.status_code == 303 and "ok=" in response.headers["location"]

    body = client.get("/api/orders").json()
    assert body["total"] == 1
    order = body["items"][0]
    assert order["pr_number"] == "PR-2026-1077"
    assert order["people"] == 2 and order["ordered"] == 18
    assert order["status"] == "awaiting_measurement"
    assert body["pending_pieces"] == 18


def test_the_order_form_demands_a_pr_number(client):
    response = client.post(
        "/orders",
        data={"employee_number": "E002", "qty_SHIRT": "2"},
        follow_redirects=False,
    )
    assert response.status_code == 303 and "PR+number" in response.headers["location"]


def test_order_form_without_a_staff_member_is_rejected(client):
    response = client.post("/orders", data={"pr_number": "PR-1", "qty_SHIRT": "2"},
                           follow_redirects=False)
    assert "err=" in response.headers["location"]


def test_recording_the_measurement_visit_over_the_api(client):
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 4}]})
    assert client.get("/api/orders/PR-1").json()["status"] == "awaiting_measurement"

    body = client.post("/api/orders/PR-1/measurement", json={}).json()
    assert body["measured"] is not None
    assert body["status"] == "pending"


def test_receiving_part_of_an_order_from_the_log(client):
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 4}]})
    client.post("/api/orders/PR-1/measurement", json={})

    response = client.post("/orders/PR-1/E002/SHIRT/deliver", data={"quantity": "1"},
                           follow_redirects=False)
    assert response.status_code == 303 and "ok=" in response.headers["location"]

    order = client.get("/api/orders/PR-1").json()
    assert order["status"] == "partially_delivered"
    assert (order["delivered"], order["outstanding"]) == (1, 3)


def test_a_delivery_before_the_measurement_is_refused(client):
    """The tailor cannot deliver clothes nobody has been measured for."""
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 2}]})
    response = client.post("/api/orders/PR-1/deliveries", json={
        "parts": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 2}]})
    assert response.status_code == 400
    assert "measurement" in response.json()["detail"]


def test_delivery_shows_up_on_the_employees_record(client):
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 1}]})
    client.post("/api/orders/PR-1/measurement", json={})
    client.post("/api/orders/PR-1/deliveries", json={
        "parts": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 1}]})

    statuses = {i["item_code"]: i["status"]
                for i in client.get("/api/employees/E002/uniforms").json()["items"]}
    assert statuses["SHIRT"] == "ok"


def test_action_needed_endpoint_ranks_renewals_above_orders(client):
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 2}]})
    rows = client.get("/api/action-needed").json()["items"]
    kinds = [r["kind"] for r in rows]
    if "order" in kinds and "renewal" in kinds:
        assert kinds.index("renewal") < kinds.index("order")


def test_dashboard_shows_pending_delivery(client):
    client.post("/api/orders", json={
        "pr_number": "PR-1",
        "lines": [{"employee_number": "E002", "item_code": "SHIRT", "quantity": 5}]})
    counts = client.get("/api/dashboard/summary").json()["counts"]
    assert counts["pending_delivery"] == 5
    assert counts["open_orders"] == 1
    assert "Pieces pending delivery" in client.get("/").text


# --- sign-in ---

@pytest.fixture
def secure_client(tmp_path, monkeypatch):
    """A client with sign-in switched on: one admin, one read-only user."""
    from app.auth import hash_password

    workbook = clean_workbook(tmp_path / "uniforms.xlsx")
    monkeypatch.setenv("WORKBOOK_PATH", str(workbook))
    monkeypatch.setenv("BACKUP_DIR", str(tmp_path / "bak"))
    monkeypatch.setenv("LEDGER_PATH", str(tmp_path / "ledger.sqlite3"))
    monkeypatch.setenv("REMINDERS_ENABLED", "0")
    monkeypatch.setenv("AUTH_ENABLED", "1")
    monkeypatch.setenv("SECRET_KEY", "test-key")
    monkeypatch.setenv(
        "AUTH_USERS",
        f"boss:{hash_password('good-pass')}:admin,desk:{hash_password('desk-pass')}:viewer",
    )

    from app import config, deps

    caches = (config.get_settings, deps.get_store, deps.get_service,
              deps.get_ledger, deps.get_auth)
    for c in caches:
        c.cache_clear()
    from app.main import app

    with TestClient(app) as c:
        yield c
    for c in caches:
        c.cache_clear()


def _sign_in(client, username, password):
    return client.post("/login", data={"username": username, "password": password},
                       follow_redirects=False)


def test_pages_redirect_to_sign_in_when_locked(secure_client):
    response = secure_client.get("/employees", follow_redirects=False)
    assert response.status_code == 303
    assert "/login" in response.headers["location"]


def test_api_returns_401_rather_than_a_redirect(secure_client):
    assert secure_client.get("/api/employees").status_code == 401


def test_login_page_is_reachable_without_signing_in(secure_client):
    assert secure_client.get("/login").status_code == 200


def test_healthz_stays_open_for_monitoring(secure_client):
    assert secure_client.get("/healthz").status_code == 200


def test_wrong_password_is_refused(secure_client):
    response = _sign_in(secure_client, "boss", "wrong")
    assert "err=" in response.headers["location"]
    assert secure_client.get("/api/employees").status_code == 401


def test_signing_in_grants_access(secure_client):
    _sign_in(secure_client, "boss", "good-pass")
    assert secure_client.get("/api/employees").status_code == 200
    assert secure_client.get("/employees").status_code == 200


def test_a_read_only_user_can_look_but_not_change(secure_client):
    _sign_in(secure_client, "desk", "desk-pass")
    assert secure_client.get("/api/employees").status_code == 200
    response = secure_client.patch("/api/employees/E002", json={"full_name": "Nope"})
    assert response.status_code == 403
    assert secure_client.get("/api/employees/E002").json()["employee"]["full_name"] != "Nope"


def test_an_admin_can_change_records(secure_client):
    _sign_in(secure_client, "boss", "good-pass")
    response = secure_client.patch("/api/employees/E002", json={"full_name": "Renamed"})
    assert response.status_code == 200


def test_edits_are_attributed_to_the_signed_in_user_not_a_typed_name(secure_client, tmp_path):
    """The whole point of sign-in: 'who' becomes a fact, not a claim."""
    _sign_in(secure_client, "boss", "good-pass")
    secure_client.post(
        "/employees/E002/edit",
        data={"full_name": "Renamed", "who": "somebody else", "reason": "typo"},
        follow_redirects=False,
    )
    secure_client.post("/api/workbook/flush")

    wb = load_workbook(tmp_path / "uniforms.xlsx")
    rows = list(wb["ChangeLog"].iter_rows(values_only=True))
    wb.close()
    header, entry = rows[0], rows[-1]
    assert entry[header.index("Who")] == "boss"


def test_signing_out_revokes_access(secure_client):
    _sign_in(secure_client, "boss", "good-pass")
    secure_client.post("/logout", follow_redirects=False)
    assert secure_client.get("/api/employees").status_code == 401


def test_a_tampered_session_cookie_is_rejected(secure_client):
    from app.auth import COOKIE

    _sign_in(secure_client, "desk", "desk-pass")
    token = secure_client.cookies.get(COOKIE)
    secure_client.cookies.set(COOKIE, token[:-6] + "AAAAAA")
    assert secure_client.get("/api/employees").status_code == 401


def test_the_api_cannot_forge_who_made_a_change(secure_client, tmp_path):
    """A caller claiming to be someone else is overruled by the session."""
    _sign_in(secure_client, "boss", "good-pass")
    secure_client.patch("/api/items/SHIRT",
                        json={"renewal_cycle_months": 18, "who": "not-me"})
    secure_client.post("/api/workbook/flush")

    wb = load_workbook(tmp_path / "uniforms.xlsx")
    rows = list(wb["ChangeLog"].iter_rows(values_only=True))
    wb.close()
    header, entry = rows[0], rows[-1]
    assert entry[header.index("Who")] == "boss"


def test_an_override_is_authorised_by_the_signed_in_user(secure_client):
    _sign_in(secure_client, "boss", "good-pass")
    body = secure_client.put(
        "/api/employees/E002/overrides/SHIRT",
        json={"next_due": "2030-01-01", "reason": "warranty replacement",
              "authorised_by": "someone-else"},
    ).json()
    assert body["authorised_by"] == "boss"
