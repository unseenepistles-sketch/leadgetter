"""Describing an edit made in Excel, before it is allowed to take effect."""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from app.domain.models import Employee, EmployeeStatus, Issuance
from app.excelstore.review import diff_snapshots
from app.excelstore.workbook import WorkbookStore
from tests.fixtures import clean_workbook


@pytest.fixture
def store(tmp_path):
    s = WorkbookStore(clean_workbook(tmp_path / "w.xlsx"), backup_dir=tmp_path / "bak")
    s.load()
    return s


def edited(store, fn):
    """Apply ``fn`` to the workbook the way a person in Excel would, then read
    it back as a candidate without adopting it."""
    wb = load_workbook(store.path)
    fn(wb)
    wb.save(store.path)
    wb.close()
    return store.read()


# --- nothing happening ---

def test_an_unchanged_file_reports_nothing(store):
    assert diff_snapshots(store.snapshot, store.read()).is_empty


def test_the_headline_says_so(store):
    assert diff_snapshots(store.snapshot, store.read()).headline() == "No changes."


# --- the cases that matter ---

def test_a_deleted_employee_is_reported_as_a_removal(store):
    """The mistake this whole gate exists for."""
    def delete_a_row(wb):
        wb["Employees"].delete_rows(2)

    diff = diff_snapshots(store.snapshot, edited(store, delete_a_row))
    removals = diff.of("employee", "removed")
    assert len(removals) == 1
    assert "E001" in removals[0].key
    assert diff.removals            # surfaced as a deletion, not buried


def test_an_added_employee_is_reported(store):
    def add_a_row(wb):
        wb["Employees"].append(
            ["E010", "New Starter", "new@x.com", "", "Ops", "Officer",
             date(2026, 1, 5), "Active", "M", "32", "40R", "9"])

    diff = diff_snapshots(store.snapshot, edited(store, add_a_row))
    added = diff.of("employee", "added")
    assert len(added) == 1
    assert "New Starter" in added[0].key


def test_an_edited_field_names_the_before_and_after(store):
    def change_a_role(wb):
        headers = [c.value for c in wb["Employees"][1]]
        column = headers.index("Role") + 1
        wb["Employees"].cell(row=2, column=column, value="Supervisor")

    diff = diff_snapshots(store.snapshot, edited(store, change_a_role))
    [change] = diff.of("employee", "edited")
    assert change.fields[0].field == "role"
    assert change.fields[0].after == "Supervisor"
    assert "Supervisor" in change.summary


def test_a_row_moved_but_not_changed_is_not_a_change(store):
    """Inserting a line above shifts every row number below it. Reporting that
    as sixty edits would bury the one that matters."""
    def insert_a_blank(wb):
        wb["Employees"].insert_rows(2)

    assert diff_snapshots(store.snapshot, edited(store, insert_a_blank)).is_empty


def test_a_new_delivery_is_reported(store):
    def add_a_handover(wb):
        wb["Issuances"].append(
            ["E002", "SHIRT", date(2024, 7, 1), 3, "M", "hand-typed", 12, ""])

    diff = diff_snapshots(store.snapshot, edited(store, add_a_handover))
    assert len(diff.of("delivery", "added")) == 1


def test_two_identical_handovers_are_two_changes(store):
    """Issuances have no natural key, so they are compared as a multiset — two
    shirts issued in two transactions on one day is a real thing."""
    def add_two(wb):
        for _ in range(2):
            wb["Issuances"].append(
                ["E002", "SHIRT", date(2024, 7, 1), 1, "M", "stores", 12, ""])

    diff = diff_snapshots(store.snapshot, edited(store, add_two))
    assert len(diff.of("delivery", "added")) == 2


def test_a_deleted_delivery_is_reported(store):
    def delete_a_handover(wb):
        wb["Issuances"].delete_rows(2)

    diff = diff_snapshots(store.snapshot, edited(store, delete_a_handover))
    assert len(diff.of("delivery", "removed")) == 1


def test_a_row_that_became_unreadable_is_called_out(store):
    def break_a_date(wb):
        headers = [c.value for c in wb["Employees"][1]]
        column = headers.index("Join Date") + 1
        wb["Employees"].cell(row=2, column=column, value="not a date")

    diff = diff_snapshots(store.snapshot, edited(store, break_a_date))
    assert diff.new_problems or diff.of("employee", "edited")


def test_a_renewal_period_change_is_reported(store):
    """Policy lives in the sheet, so this is a legitimate edit — but she should
    still see that it happened."""
    def change_the_cycle(wb):
        headers = [c.value for c in wb["UniformItems"][1]]
        column = headers.index("Renewal Cycle Months") + 1
        wb["UniformItems"].cell(row=2, column=column, value=18)

    diff = diff_snapshots(store.snapshot, edited(store, change_the_cycle))
    [change] = diff.of("item", "edited")
    assert change.fields[0].after == 18


# --- how it reads ---

def test_the_headline_counts_each_kind(store):
    def several_things(wb):
        wb["Employees"].delete_rows(2)
        wb["Issuances"].append(
            ["E002", "SHIRT", date(2024, 7, 1), 3, "M", "stores", 12, ""])

    diff = diff_snapshots(store.snapshot, edited(store, several_things))
    headline = diff.headline()
    assert "employees removed" in headline
    assert "deliverys added" in headline or "delivery" in headline


def test_dates_and_blanks_read_as_words_not_repr(store):
    def clear_an_email(wb):
        headers = [c.value for c in wb["Employees"][1]]
        column = headers.index("Email") + 1
        # Assign through .value: ws.cell(..., value=None) means "leave it alone".
        wb["Employees"].cell(row=2, column=column).value = None

    diff = diff_snapshots(store.snapshot, edited(store, clear_an_email))
    [change] = diff.of("employee", "edited")
    assert "(blank)" in change.summary


# --- what the hold protects ---

def test_a_write_made_during_the_hold_survives_and_lands_after(store, tmp_path):
    """She carries on recording while an edit waits. Her work must not be lost,
    and must not be written over somebody's unreviewed change either."""
    from app.service import UniformService

    service = UniformService(store)
    service.place_order("PR-1", [{"employee_number": "E002", "item_code": "SHIRT",
                                  "quantity": 2}], ordered_date=date(2026, 1, 5))
    store.flush()

    edited(store, lambda wb: wb["Employees"].delete_rows(2))
    assert store.reload_if_changed() is True

    service.edit_order_header("PR-1", {"supplier_ref": "Bespoke Workwear"})
    assert store.pending_writes > 0
    assert store.flush() == 0                      # held, not written
    assert store.pending_writes > 0                # and not dropped

    store.accept_pending()
    assert store.flush() > 0

    ws = load_workbook(store.path)["Orders"]
    headers = [c.value for c in ws[1]]
    tailors = {r[headers.index("Tailor")]
               for r in ws.iter_rows(min_row=2, values_only=True)}
    assert tailors == {"Bespoke Workwear"}


def test_the_service_reports_the_change_for_review(store):
    from app.service import UniformService

    service = UniformService(store)
    assert service.pending_workbook_change() is None

    edited(store, lambda wb: wb["Employees"].delete_rows(2))
    store.reload_if_changed()

    diff = service.pending_workbook_change()
    assert diff is not None and diff.removals
    service.accept_workbook_change()
    assert service.pending_workbook_change() is None


def test_shifting_rows_does_not_report_old_problems_as_new(store):
    """Deleting a row renumbers everything below it. Keying problems on the row
    number would report every existing one as brand new."""
    diff = diff_snapshots(store.snapshot,
                          edited(store, lambda wb: wb["Employees"].delete_rows(2)))
    assert diff.new_problems == []
    assert diff.of("employee", "removed")     # the real change is still reported
