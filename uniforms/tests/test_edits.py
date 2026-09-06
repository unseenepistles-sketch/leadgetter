"""Editing the system of record: what changes, what is refused, what is logged."""
from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from app.domain.models import EmployeeStatus
from app.excelstore.workbook import WorkbookStore
from app.service import NotFound, UniformService, ValidationError
from tests.fixtures import clean_workbook


@pytest.fixture
def service(tmp_path):
    store = WorkbookStore(clean_workbook(tmp_path / "w.xlsx"), backup_dir=tmp_path / "bak")
    store.load()
    return UniformService(store)


def _sheet(service, name):
    wb = load_workbook(service.store.path)
    rows = list(wb[name].iter_rows(values_only=True))
    wb.close()
    return rows


# --- editing employees ---

def test_edit_employee_details(service):
    updated = service.edit_employee("E002", {"full_name": "Amina Y. Bello",
                                             "department": "Stations"}, who="njiru")
    assert updated.full_name == "Amina Y. Bello"
    assert updated.department == "Stations"
    assert service.employee("E002").full_name == "Amina Y. Bello"


def test_edit_employee_reaches_the_workbook(service):
    service.edit_employee("E002", {"full_name": "Amina Y. Bello"}, who="njiru")
    service.store.flush()
    rows = _sheet(service, "Employees")
    header = rows[0]
    row = next(r for r in rows[1:] if r[header.index("Employee Number")] == "E002")
    assert row[header.index("Full Name")] == "Amina Y. Bello"


def test_editing_sizes_updates_the_profile(service):
    updated = service.edit_employee("E002", {"shirt_size": "XL"}, who="njiru")
    assert updated.sizes["shirt"] == "XL"
    # and the next issue picks it up
    issued = service.record_issuance("E002", "SHIRT")
    assert issued.size == "XL"


def test_editing_the_role_changes_what_they_are_entitled_to(service):
    before = {s.item_code for s in service.statuses_for("E002")}
    service.edit_employee("E002", {"role": "clerk"}, who="njiru")
    after = {s.item_code for s in service.statuses_for("E002")}
    assert before != after


def test_editing_a_join_date_moves_the_renewal_clock(service):
    service.edit_employee("E002", {"join_date": "2024-01-15"}, who="njiru")
    assert service.employee("E002").join_date == date(2024, 1, 15)


def test_marking_an_employee_as_left_removes_them_from_reports(service):
    service.edit_employee("E002", {"status": "left"}, who="njiru")
    assert service.employee("E002").status is EmployeeStatus.LEFT
    assert all(s.employee_number != "E002" for s in service.all_statuses())


def test_an_unknown_status_is_refused(service):
    with pytest.raises(ValidationError, match="status must be one of"):
        service.edit_employee("E002", {"status": "on holiday"}, who="njiru")


def test_unknown_fields_are_refused(service):
    with pytest.raises(ValidationError, match="cannot edit"):
        service.edit_employee("E002", {"salary": 1}, who="njiru")


def test_editing_nothing_is_a_no_op(service):
    before = service.employee("E002")
    assert service.edit_employee("E002", {"full_name": before.full_name}, who="njiru") == before
    assert service.store.pending_writes == 0


# --- editing orders ---

def raised(service, qty, pr="PR-1"):
    """One line, measured, ready to be edited or delivered against."""
    service.place_order(pr, [{"employee_number": "E002", "item_code": "SHIRT",
                              "quantity": qty}], ordered_date=date(2026, 1, 5))
    service.record_measurement(pr, date(2026, 1, 10))
    return pr


def test_edit_order_quantity_and_remarks(service):
    pr = raised(service, 2)
    line = service.edit_order(pr, "E002", "SHIRT",
                              {"quantity": 5, "notes": "size change"}, who="njiru")
    assert line.order.quantity == 5 and line.order.notes == "size change"
    assert line.pending == 5


def test_order_quantity_cannot_drop_below_what_arrived(service):
    pr = raised(service, 4)
    service.receive_delivery(pr, "E002", "SHIRT", quantity=3)
    with pytest.raises(ValidationError, match="already delivered"):
        service.edit_order(pr, "E002", "SHIRT", {"quantity": 2}, who="njiru")


def test_reducing_quantity_to_what_arrived_closes_the_order(service):
    pr = raised(service, 4)
    service.receive_delivery(pr, "E002", "SHIRT", quantity=2)
    line = service.edit_order(pr, "E002", "SHIRT", {"quantity": 2}, who="njiru")
    assert line.status.value == "delivered" and line.pending == 0


def test_cancelling_an_order(service):
    pr = raised(service, 2)
    line = service.edit_order(pr, "E002", "SHIRT", {"cancelled": True}, who="njiru")
    assert line.status.value == "cancelled"
    assert service.pending_delivery() == 0


def test_editing_a_line_leaves_its_neighbours_alone(service):
    """The (PR, person, item) key has to isolate one row out of a bulk order."""
    service.place_order("PR-1", [
        {"employee_number": "E002", "item_code": "SHIRT", "quantity": 2},
        {"employee_number": "E003", "item_code": "SHIRT", "quantity": 2},
    ], ordered_date=date(2026, 1, 5))
    service.edit_order("PR-1", "E002", "SHIRT", {"quantity": 7}, who="njiru")
    assert service.order_line("PR-1", "E002", "SHIRT").order.quantity == 7
    assert service.order_line("PR-1", "E003", "SHIRT").order.quantity == 2


# --- editing deliveries ---

def test_edit_a_delivery_date(service):
    pr = raised(service, 1)
    service.receive_delivery(pr, "E002", "SHIRT")
    service.store.flush()
    service.store.load()

    row = service.snapshot.issuances_for("E002")[0].row
    updated = service.edit_delivery("E002", row, {"issued_date": "2026-01-15"}, who="njiru")
    assert updated.issued_date == date(2026, 1, 15)
    # the renewal clock follows the corrected date
    assert service.order_line(pr, "E002", "SHIRT").next_due == date(2027, 1, 15)
