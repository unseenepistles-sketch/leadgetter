"""Builders for test workbooks.

``messy_workbook`` deliberately reproduces what real spreadsheets look like:
renamed and reordered columns, four date formats, blank spacer rows, duplicate
keys and junk values. A fixture that is too tidy lets bugs through to production.
"""
from __future__ import annotations

import random
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook


def clean_workbook(path: Path) -> Path:
    wb = Workbook()
    wb.remove(wb.active)

    ws = wb.create_sheet("Employees")
    ws.append(["Employee Number", "Full Name", "Email", "Manager Email", "Department",
               "Role", "Join Date", "Status", "Shirt Size", "Trouser Size", "Blazer Size"])
    ws.append(["E001", "Sam Okoro", "sam@co.test", "boss@co.test", "Ops",
               "officer", date(2020, 1, 15), "Active", "L", "34", "42R"])
    ws.append(["E002", "Amina Yusuf", "amina@co.test", "boss@co.test", "Ops",
               "officer", date(2023, 6, 1), "Active", "M", "30", "38R"])
    ws.append(["E003", "Tunde Bello", "tunde@co.test", "boss@co.test", "Admin",
               "clerk", date(2019, 3, 10), "Active", "XL", "36", "44R"])
    ws.append(["E004", "Grace Mwangi", "grace@co.test", "boss@co.test", "Ops",
               "officer", date(2021, 9, 1), "Left", "S", "28", "36R"])

    ws = wb.create_sheet("UniformItems")
    ws.append(["Item Code", "Name", "Category", "Renewal Cycle Months",
               "Default Quantity", "Size Key", "Active"])
    ws.append(["SHIRT", "Long-sleeve shirt", "Shirt", 12, 3, "shirt", "Yes"])
    ws.append(["TROUSER", "Work trousers", "Trouser", 12, 2, "trouser", "Yes"])
    ws.append(["BLAZER", "Blazer", "Blazer", 24, 1, "blazer", "Yes"])
    ws.append(["BADGE", "Old name badge", "Accessory", 12, 1, "", "No"])

    ws = wb.create_sheet("Entitlements")
    ws.append(["Role", "Item Code", "Quantity"])
    for role in ("officer", "clerk"):
        ws.append([role, "SHIRT", 3])
        ws.append([role, "TROUSER", 2])
    ws.append(["officer", "BLAZER", 1])

    ws = wb.create_sheet("Issuances")
    ws.append(["Employee Number", "Item Code", "Issued Date", "Quantity", "Size",
               "Issued By", "Cycle Months", "Notes"])
    ws.append(["E001", "SHIRT", date(2024, 1, 15), 3, "L", "stores", 12, ""])
    ws.append(["E001", "TROUSER", date(2024, 1, 15), 2, "34", "stores", 12, ""])
    ws.append(["E001", "BLAZER", date(2023, 1, 15), 1, "42R", "stores", 24, ""])
    ws.append(["E003", "SHIRT", date(2019, 3, 12), 3, "XL", "stores", 12, ""])

    wb.save(path)
    wb.close()
    return path


def messy_workbook(path: Path) -> Path:
    """Everything that goes wrong in a real spreadsheet, in one file."""
    wb = Workbook()
    wb.remove(wb.active)

    # Tab named "Staff List", columns reordered and renamed, mixed date formats.
    ws = wb.create_sheet("Staff List")
    ws.append(["Name", "Dept", "Staff ID", "Job Title", "Date Joined", "Email", "Status", "Shirt"])
    ws.append(["Sam Okoro", "Ops", "E001", "officer", "15/01/2020", "sam@co.test", "Active", "L"])
    ws.append([None, None, None, None, None, None, None, None])          # spacer row
    ws.append(["Amina Yusuf", "Ops", "E002", "officer", datetime(2023, 6, 1), "amina@co.test", "ACTIVE", "M"])
    ws.append(["Tunde Bello", "Admin", "E003", "clerk", "10-Mar-2019", "tunde@co.test", "y", "XL"])
    ws.append(["Grace Mwangi", "Ops", "E004", "officer", 44440, "grace@co.test", "Left", "S"])  # Excel serial
    ws.append(["Duplicate Sam", "Ops", "E001", "officer", "01/01/2021", "dup@co.test", "Active", "L"])
    ws.append(["Broken Date", "Ops", "E005", "officer", "13/13/2020", "bad@co.test", "Active", "M"])
    ws.append(["No Join Date", "Ops", "E006", "officer", None, "nojoin@co.test", "Active", "M"])
    ws.append([None, "Ops", None, "officer", "01/01/2020", "noid@co.test", "Active", "M"])      # no id

    ws = wb.create_sheet("Uniforms")
    ws.append(["Code", "Description", "Type", "Cycle", "Qty", "In Use"])
    ws.append(["SHIRT", "Long-sleeve shirt", "Shirt", "12", 3, "yes"])
    ws.append(["TROUSER", "Work trousers", "Trouser", 12, "2", "TRUE"])
    ws.append(["BLAZER", "Blazer", "Blazer", "24 months", 1, 1])
    ws.append(["BADGE", "Old badge", "Accessory", 0, 1, "no"])           # zero cycle

    ws = wb.create_sheet("Allocations")
    ws.append(["Grade", "Item", "Number"])
    ws.append(["officer", "SHIRT", 3])
    ws.append(["officer", "TROUSER", 2])
    ws.append(["officer", "BLAZER", 1])
    ws.append(["clerk", "SHIRT", 3])
    ws.append(["clerk", "GHOST", 1])                                      # unknown item
    ws.append(["", "SHIRT", 1])                                           # no role

    ws = wb.create_sheet("Issue Log")
    ws.append(["Staff ID", "Item", "Date Issued", "Qty", "Size", "Given By", "Remarks"])
    ws.append(["E001", "SHIRT", "15/01/2024", 3, "L", "stores", ""])
    ws.append(["E001", "TROUSER", "15-01-2024", "2", "34", "stores", ""])
    ws.append(["E001", "BLAZER", datetime(2023, 1, 15), 1, "42R", "stores", ""])
    ws.append(["E003", "SHIRT", "12 March 2019", 3, "XL", "stores", ""])
    ws.append(["E003", "SHIRT", "2019-03-12", 3, "XXL", "stores", "same day, later row"])
    ws.append(["E002", "SHIRT", "not a date", 3, "M", "stores", ""])      # unreadable
    ws.append([None, "SHIRT", "01/01/2024", 3, "M", "stores", ""])        # no staff id
    ws.append([None, None, None, None, None, None, None])

    wb.save(path)
    wb.close()
    return path


def large_workbook(path: Path, employees: int = 10_000, issuances_each: int = 6) -> Path:
    """A workbook at the client's real scale, for load and timing checks."""
    rng = random.Random(42)
    wb = Workbook(write_only=True)

    ws = wb.create_sheet("Employees")
    ws.append(["Employee Number", "Full Name", "Email", "Manager Email", "Department",
               "Role", "Join Date", "Status", "Shirt Size", "Trouser Size", "Blazer Size"])
    departments = ["Ops", "Admin", "Field", "Depot", "HQ"]
    for i in range(1, employees + 1):
        ws.append([
            f"E{i:06d}", f"Employee {i}", f"e{i}@co.test", "boss@co.test",
            departments[i % len(departments)], "officer",
            date(2015, 1, 1) + timedelta(days=rng.randint(0, 3500)),
            "Active", "L", "34", "42R",
        ])

    ws = wb.create_sheet("UniformItems")
    ws.append(["Item Code", "Name", "Category", "Renewal Cycle Months",
               "Default Quantity", "Size Key", "Active"])
    ws.append(["SHIRT", "Shirt", "Shirt", 12, 3, "shirt", "Yes"])
    ws.append(["TROUSER", "Trousers", "Trouser", 12, 2, "trouser", "Yes"])
    ws.append(["BLAZER", "Blazer", "Blazer", 24, 1, "blazer", "Yes"])

    ws = wb.create_sheet("Entitlements")
    ws.append(["Role", "Item Code", "Quantity"])
    ws.append(["officer", "SHIRT", 3])
    ws.append(["officer", "TROUSER", 2])
    ws.append(["officer", "BLAZER", 1])

    ws = wb.create_sheet("Issuances")
    ws.append(["Employee Number", "Item Code", "Issued Date", "Quantity", "Size",
               "Issued By", "Cycle Months", "Notes"])
    codes = ["SHIRT", "TROUSER", "BLAZER"]
    for i in range(1, employees + 1):
        for n in range(issuances_each):
            ws.append([
                f"E{i:06d}", codes[n % 3],
                date(2019, 1, 1) + timedelta(days=rng.randint(0, 2200)),
                1, "L", "stores", 12 if n % 3 != 2 else 24, "",
            ])

    wb.save(path)
    wb.close()
    return path
