from __future__ import annotations

from datetime import date

import pytest
from openpyxl import load_workbook

from app.domain.models import EmployeeStatus, Issuance
from app.excelstore.workbook import WorkbookStore, create_blank_workbook
from tests.fixtures import clean_workbook, messy_workbook


@pytest.fixture
def clean(tmp_path):
    store = WorkbookStore(clean_workbook(tmp_path / "clean.xlsx"), backup_dir=tmp_path / "bak")
    store.load()
    return store


@pytest.fixture
def messy(tmp_path):
    store = WorkbookStore(messy_workbook(tmp_path / "messy.xlsx"), backup_dir=tmp_path / "bak")
    store.load()
    return store


# --- reading a well-formed workbook ---

def test_reads_all_four_sheets(clean):
    counts = clean.snapshot.counts
    assert counts["employees"] == 4
    assert counts["items"] == 4
    assert counts["issuances"] == 4


def test_indexes_issuances_by_employee(clean):
    assert len(clean.snapshot.issuances_for("E001")) == 3
    assert clean.snapshot.issuances_for("E999") == []


def test_entitlements_resolve_by_role_case_insensitively(clean):
    assert len(clean.snapshot.entitlements_for("officer")) == 3
    assert len(clean.snapshot.entitlements_for("OFFICER")) == 3
    assert len(clean.snapshot.entitlements_for("clerk")) == 2


def test_left_employees_are_marked(clean):
    assert clean.snapshot.employees["E004"].status is EmployeeStatus.LEFT
    assert not clean.snapshot.employees["E004"].is_active


# --- surviving a spreadsheet written by humans ---

def test_finds_sheets_under_their_real_world_names(messy):
    """Tabs are 'Staff List', 'Uniforms', 'Allocations', 'Issue Log'."""
    counts = messy.snapshot.counts
    assert counts["employees"] >= 5
    assert counts["items"] == 4


def test_columns_matched_by_name_not_position(messy):
    """Staff ID is the third column here, not the first."""
    assert messy.snapshot.employees["E001"].full_name == "Sam Okoro"
    assert messy.snapshot.employees["E001"].department == "Ops"


@pytest.mark.parametrize(
    "emp_no,expected",
    [
        ("E001", date(2020, 1, 15)),   # "15/01/2020" text, day-first
        ("E002", date(2023, 6, 1)),    # real datetime
        ("E003", date(2019, 3, 10)),   # "10-Mar-2019"
        ("E004", date(2021, 9, 1)),    # Excel serial 44440
    ],
)
def test_parses_every_date_format_in_the_sheet(messy, emp_no, expected):
    assert messy.snapshot.employees[emp_no].join_date == expected


def test_blank_spacer_rows_are_skipped(messy):
    assert all(e.employee_number for e in messy.snapshot.employees.values())


def test_duplicate_employee_number_is_reported_not_silently_overwritten(messy):
    assert messy.snapshot.employees["E001"].full_name == "Sam Okoro"
    assert any("duplicate" in p.message for p in messy.snapshot.problems)


def test_unreadable_join_date_flags_the_employee(messy):
    emp = messy.snapshot.employees["E005"]
    assert emp.problems and "13/13/2020" in emp.problems[0]


def test_row_with_no_employee_number_is_reported(messy):
    assert any("no employee number" in p.message for p in messy.snapshot.problems)


def test_unreadable_issue_date_drops_the_row_loudly(messy):
    """E002's only issuance has a junk date — better absent than wrong."""
    assert messy.snapshot.issuances_for("E002") == []
    assert any("unreadable issue date" in p.message for p in messy.snapshot.problems)


def test_cycle_written_as_text_is_understood(messy):
    assert messy.snapshot.items["SHIRT"].renewal_cycle_months == 12
    assert messy.snapshot.items["BLAZER"].renewal_cycle_months == 24  # "24 months"


def test_zero_cycle_falls_back_to_twelve_and_is_reported(messy):
    assert messy.snapshot.items["BADGE"].renewal_cycle_months == 12
    assert any("cycle must be positive" in p.message for p in messy.snapshot.problems)


def test_active_flag_understood_in_all_its_spellings(messy):
    assert messy.snapshot.items["SHIRT"].active is True
    assert messy.snapshot.items["TROUSER"].active is True
    assert messy.snapshot.items["BLAZER"].active is True
    assert messy.snapshot.items["BADGE"].active is False


def test_entitlement_pointing_at_unknown_item_is_reported(messy):
    codes = {e.item_code for e in messy.snapshot.entitlements_for("clerk")}
    assert "GHOST" not in codes
    assert any("unknown item code" in p.message for p in messy.snapshot.problems)


# --- writing ---

def test_appended_issuance_is_visible_immediately(clean):
    clean.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1), 3, "M", "stores", 12))
    assert len(clean.snapshot.issuances_for("E002")) == 1
    assert clean.pending_writes == 1


def test_flush_persists_to_the_file(clean):
    clean.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1), 3, "M", "stores", 12))
    assert clean.flush() == 1
    assert clean.pending_writes == 0

    wb = load_workbook(clean.path)
    rows = list(wb["Issuances"].iter_rows(min_row=2, values_only=True))
    wb.close()
    assert len(rows) == 5
    assert rows[-1][0] == "E002"


def test_flushed_rows_survive_a_reload(clean):
    clean.append_issuance(Issuance("E002", "BLAZER", date(2024, 6, 1), 1, "38R", "stores", 24))
    clean.flush()
    clean.load()
    assert len(clean.snapshot.issuances_for("E002")) == 1


def test_append_respects_the_sheets_own_column_order(messy):
    """The 'Issue Log' tab has no Cycle Months column and Remarks last."""
    messy.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1), 3, "M", "stores", 12, "note"))
    messy.flush()
    wb = load_workbook(messy.path)
    header = next(wb["Issue Log"].iter_rows(min_row=1, max_row=1, values_only=True))
    last = list(wb["Issue Log"].iter_rows(values_only=True))[-1]
    wb.close()
    assert last[header.index("Staff ID")] == "E002"
    assert last[header.index("Size")] == "M"
    assert last[header.index("Remarks")] == "note"


def test_flush_backs_up_first(clean):
    clean.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1)))
    clean.flush()
    assert list(clean.backup_dir.glob("*.xlsx"))


def test_backups_are_capped(tmp_path):
    store = WorkbookStore(
        clean_workbook(tmp_path / "w.xlsx"), backup_dir=tmp_path / "bak", max_backups=3
    )
    store.load()
    for i in range(6):
        store.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1 + i)))
        store.flush()
    assert len(list(store.backup_dir.glob("*.xlsx"))) <= 3


def test_flush_with_nothing_queued_is_a_no_op(clean):
    assert clean.flush() == 0


def test_write_leaves_no_temp_file_behind(clean):
    clean.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1)))
    clean.flush()
    assert not list(clean.path.parent.glob(".*.writing.xlsx"))


def test_bulk_append_writes_in_one_pass(clean):
    rows = [Issuance(f"E00{i}", "SHIRT", date(2024, 6, 1)) for i in (1, 2, 3)]
    clean.append_issuances(rows)
    assert clean.flush() == 3


# --- reload when a human edits the file underneath us ---

def test_reload_if_changed_is_a_no_op_when_untouched(clean):
    assert clean.reload_if_changed() is False


def test_reload_if_changed_picks_up_an_external_edit(clean):
    wb = load_workbook(clean.path)
    wb["Issuances"].append(["E002", "SHIRT", date(2024, 7, 1), 3, "M", "hand-typed", 12, ""])
    wb.save(clean.path)
    wb.close()
    assert clean.reload_if_changed() is True
    assert len(clean.snapshot.issuances_for("E002")) == 1


def test_our_own_write_does_not_trigger_a_reload(clean):
    clean.append_issuance(Issuance("E002", "SHIRT", date(2024, 6, 1)))
    clean.flush()
    assert clean.reload_if_changed() is False


# --- bootstrapping ---

def test_create_blank_workbook_is_immediately_loadable(tmp_path):
    path = create_blank_workbook(tmp_path / "new.xlsx")
    store = WorkbookStore(path, backup_dir=tmp_path / "bak")
    snap = store.load()
    assert snap.counts["employees"] == 0
    assert snap.problems == []


def test_missing_join_date_is_reported_without_forcing_needs_review(tmp_path):
    """Actionable for the client, but someone with issuance history is still judgeable."""
    from openpyxl import load_workbook as _lw
    path = clean_workbook(tmp_path / "w.xlsx")
    wb = _lw(path)
    wb["Employees"].cell(row=2, column=7).value = None  # E001 loses its join date
    wb.save(path)
    wb.close()

    store = WorkbookStore(path, backup_dir=tmp_path / "bak")
    store.load()
    assert store.snapshot.employees["E001"].join_date is None
    assert store.snapshot.employees["E001"].problems == ()
    assert any("no join date" in p.message for p in store.snapshot.problems)
