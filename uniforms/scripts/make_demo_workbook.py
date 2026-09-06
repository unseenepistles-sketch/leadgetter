#!/usr/bin/env python3
"""Build a realistic demo workbook for a fictional company.

Deliberately spread across every status the system can report — up to date,
due soon, due, overdue, never issued, and a couple of rows with bad data — so a
demo shows what the tool actually does rather than an empty dashboard.

    python scripts/make_demo_workbook.py data/uniforms.xlsx
"""
from __future__ import annotations

import argparse
import random
import sys
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

HEADER_FILL = PatternFill("solid", fgColor="1F5FBF")
HEADER_FONT = Font(color="FFFFFF", bold=True)

FIRST = [
    "Amina", "Tunde", "Grace", "Samuel", "Chidi", "Fatima", "Joseph", "Ngozi", "David",
    "Halima", "Emeka", "Rita", "Musa", "Blessing", "Peter", "Zainab", "Ibrahim", "Chioma",
    "Daniel", "Aisha", "Kelechi", "Yusuf", "Adaeze", "Bello", "Esther", "Segun", "Mary",
    "Abdul", "Ifeoma", "John", "Salma", "Victor", "Hauwa", "Paul", "Nkechi", "Ahmed",
    "Ruth", "Femi", "Ladi", "Charles", "Bisi", "Umar", "Joy", "Kunle", "Sarah", "Idris",
    "Tope", "Comfort",
]
LAST = [
    "Okoro", "Bello", "Mwangi", "Adeyemi", "Nwosu", "Yusuf", "Eze", "Abubakar", "Ojo",
    "Danladi", "Okafor", "Balogun", "Sani", "Ibrahim", "Ogunleye", "Musa", "Chukwu",
    "Lawal", "Adamu", "Onyeka", "Garba", "Uche", "Suleiman", "Afolabi", "Idowu",
]

# Mirrors the operations department this is being built for: a tram network,
# where the uniform is customer-facing and the renewal cycle is a real policy.
# The three staff categories named in the requirements, nothing else.
DEPARTMENTS = {
    "OMCC": ["OMCC Controller"],
    "Operations": ["Line Supervisor"],
    "Drivers": ["Tram Driver"],
}

# The catalogue and cycles the department already uses. Shirts and trousers wear
# out yearly; structured and hard-wearing items run two years.
ITEMS = [
    # code, name, category, cycle months, qty, size key
    ("SHIRT", "Shirt", "Shirt", 12, 3, "shirt"),
    ("TROUSER", "Trousers", "Trouser", 12, 2, "trouser"),
    ("JACKET", "Jacket", "Blazer", 24, 1, "blazer"),
    ("WAISTCOAT", "Waist Coat", "Blazer", 24, 1, "blazer"),
    ("WINTERJKT", "Winter Jacket", "Blazer", 24, 1, "blazer"),
    ("SHOES", "Safety Shoes", "Shoe", 24, 1, "shoe"),
    ("TIE", "Tie", "Accessory", 24, 1, None),
    ("BELT", "Belt", "Accessory", 24, 1, None),
    ("BADGE", "Name Badge", "Accessory", 24, 1, None),
]

# Which kit each role gets. Deliberately varied so the report is not uniform.
ENTITLEMENTS = {
    "OMCC Controller": [("SHIRT", 3), ("TROUSER", 2), ("JACKET", 1), ("WAISTCOAT", 1),
                        ("WINTERJKT", 1), ("SHOES", 1), ("TIE", 1), ("BELT", 1), ("BADGE", 1)],
    "Line Supervisor": [("SHIRT", 3), ("TROUSER", 2), ("JACKET", 1), ("WAISTCOAT", 1),
                        ("WINTERJKT", 1), ("SHOES", 1), ("TIE", 1), ("BELT", 1), ("BADGE", 1)],
    "Tram Driver": [("SHIRT", 3), ("TROUSER", 2), ("JACKET", 1), ("WAISTCOAT", 1),
                    ("SHOES", 1), ("TIE", 1), ("BELT", 1), ("BADGE", 1)],
}

SHIRT_SIZES = ["S", "M", "L", "XL", "XXL"]
TROUSER_SIZES = ["28", "30", "32", "34", "36", "38"]
BLAZER_SIZES = ["36R", "38R", "40R", "42R", "44R"]
SHOE_SIZES = ["6", "7", "8", "9", "10", "11", "12"]


def _style(ws, widths: list[int]) -> None:
    for i, width in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = width
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center")
    ws.freeze_panes = "A2"


def build(path: Path, *, count: int = 48, today: date | None = None, seed: int = 7) -> Path:
    rng = random.Random(seed)
    today = today or date.today()
    wb = Workbook()
    wb.remove(wb.active)

    # ---------------------------------------------------------------- employees
    ws = wb.create_sheet("Employees")
    ws.append(["Employee Number", "Full Name", "Email", "Manager Email", "Department",
               "Role", "Join Date", "Status", "Shirt Size", "Trouser Size",
               "Blazer Size", "Shoe Size"])

    people: list[dict] = []
    used_names: set[str] = set()
    for i in range(1, count + 1):
        while True:
            name = f"{rng.choice(FIRST)} {rng.choice(LAST)}"
            if name not in used_names:
                used_names.add(name)
                break
        department = rng.choice(list(DEPARTMENTS))
        role = rng.choice(DEPARTMENTS[department])
        # Spread joins over eight years so cycles land at different points.
        joined = today - timedelta(days=rng.randint(20, 2900))
        status = "Active"
        if i % 17 == 0:
            status = "Left"
        elif i % 23 == 0:
            status = "On Leave"

        person = {
            "number": f"EMP{i:04d}",
            "name": name,
            "email": f"{name.split()[0].lower()}.{name.split()[1].lower()}@tramline-demo.test",
            "manager": f"{department.split()[0].lower()}.manager@tramline-demo.test",
            "department": department,
            "role": role,
            "joined": joined,
            "status": status,
            "shirt": rng.choice(SHIRT_SIZES),
            "trouser": rng.choice(TROUSER_SIZES),
            "blazer": rng.choice(BLAZER_SIZES),
            "shoe": rng.choice(SHOE_SIZES),
        }
        people.append(person)

    # Two rows the parser cannot fully trust, so the Data Quality page is not empty.
    people[6]["joined"] = None                       # missing join date
    people[11]["joined_raw"] = "31/02/2021"          # a date that does not exist

    for p in people:
        ws.append([
            p["number"], p["name"], p["email"], p["manager"], p["department"], p["role"],
            p.get("joined_raw", p["joined"]), p["status"],
            p["shirt"], p["trouser"], p["blazer"], p["shoe"],
        ])
    _style(ws, [17, 22, 34, 32, 14, 21, 12, 10, 10, 12, 11, 10])

    # -------------------------------------------------------------------- items
    ws = wb.create_sheet("UniformItems")
    ws.append(["Item Code", "Name", "Category", "Renewal Cycle Months",
               "Default Quantity", "Size Key", "Active"])
    for code, name, category, cycle, qty, size_key in ITEMS:
        ws.append([code, name, category, cycle, qty, size_key, "Yes"])
    _style(ws, [12, 22, 12, 22, 17, 11, 9])

    # ------------------------------------------------------------- entitlements
    ws = wb.create_sheet("Entitlements")
    ws.append(["Role", "Item Code", "Quantity"])
    for role, rows in ENTITLEMENTS.items():
        for code, qty in rows:
            ws.append([role, code, qty])
    _style(ws, [24, 12, 11])

    # ---------------------------------------------------------------- issuances
    ws = wb.create_sheet("Issuances")
    ws.append(["Employee Number", "Item Code", "Issued Date", "Quantity", "Size",
               "Issued By", "Cycle Months", "Notes"])

    cycles = {code: cycle for code, _, _, cycle, _, _ in ITEMS}
    size_keys = {code: key for code, _, _, _, _, key in ITEMS}
    issuers = ["A. Store", "B. Ojo", "C. Danladi", "Stores desk"]
    rows = 0

    for index, p in enumerate(people):
        if p["status"] == "Left":
            continue
        joined = p["joined"]
        if joined is None:
            continue
        entitled = ENTITLEMENTS[p["role"]]

        # One in six was never issued anything at all — the "missed" case that the
        # client currently has no way of finding.
        if index % 6 == 2:
            continue

        for code, qty in entitled:
            # And some individual items were skipped even where others were issued.
            if rng.random() < 0.18:
                continue
            cycle = cycles[code]
            # Pick how far through the cycle they are, so the report spreads across
            # up-to-date, due soon, due and overdue.
            roll = rng.random()
            if roll < 0.35:
                age_days = rng.randint(0, int(cycle * 30.4 * 0.5))      # comfortable
            elif roll < 0.60:
                age_days = rng.randint(int(cycle * 30.4 * 0.5), int(cycle * 30.4 * 0.98))
            elif roll < 0.80:
                age_days = rng.randint(int(cycle * 30.4 * 0.98), int(cycle * 30.4 * 1.1))
            else:
                age_days = rng.randint(int(cycle * 30.4 * 1.1), int(cycle * 30.4 * 1.8))

            issued = today - timedelta(days=age_days)
            if issued < joined:
                issued = joined
            ws.append([
                p["number"], code, issued, qty,
                p[size_keys[code]] if size_keys.get(code) else None,
                rng.choice(issuers), cycle, "",
            ])
            rows += 1

    _style(ws, [17, 12, 13, 10, 9, 14, 14, 20])

    # ------------------------------------------------------------------- orders
    # Real orders are bulk ones: one PR number covering many people and many
    # items, moving through raised -> measured -> part-delivered -> delivered.
    ws = wb.create_sheet("Orders")
    ws.append(["PR Number", "Employee Number", "Item Code", "Ordered Date", "Quantity",
               "Size", "Measured Date", "Tailor", "Cancelled", "Notes"])
    deliveries: list[list] = []
    active = [p for p in people if p["status"] == "Active" and p["joined"]]
    tailors = ["Al Noor Tailors", "Crown Uniforms", "Bespoke Workwear"]
    orders = 0

    # Four PRs at the four stages of the journey, so the log shows the spread.
    for n, (days_ago, stage) in enumerate(
        [(96, "delivered"), (61, "partial"), (28, "awaiting"), (9, "measuring")], start=1
    ):
        raised = today - timedelta(days=days_ago)
        measured = None if stage == "measuring" else raised + timedelta(days=rng.randint(6, 16))
        pr = f"PR-{raised:%Y}-{1070 + n * 13}"
        tailor = rng.choice(tailors)
        orders += 1

        for p in rng.sample(active, k=min(rng.randint(6, 12), len(active))):
            for code, qty in ENTITLEMENTS[p["role"]][: rng.randint(2, 4)]:
                qty = max(1, qty)
                ws.append([pr, p["number"], code, raised, qty,
                           p[size_keys[code]] if size_keys.get(code) else None,
                           measured, tailor, None, ""])
                if stage == "measuring" or stage == "awaiting":
                    continue
                # Partly delivered orders leave roughly a third still to come.
                got = qty if stage == "delivered" or rng.random() > 0.35 else 0
                if not got:
                    continue
                arrived = min(today, measured + timedelta(days=rng.randint(12, 34)))
                deliveries.append([
                    p["number"], code, arrived, got,
                    p[size_keys[code]] if size_keys.get(code) else None,
                    rng.choice(issuers), cycles[code], pr, "",
                ])

    _style(ws, [16, 17, 12, 13, 10, 9, 14, 20, 11, 20])

    # Deliveries are ordinary issuance rows carrying the order they fulfil, so the
    # renewal clock runs from arrival and "1/2 delivered" is always computed.
    issues = wb["Issuances"]
    issues.cell(row=1, column=8, value="Order ID")
    issues.cell(row=1, column=9, value="Notes")
    for row in deliveries:
        issues.append(row)
    rows += len(deliveries)

    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    wb.close()
    print(f"wrote {path}: {len(people)} employees, {len(ITEMS)} items, "
          f"{rows} issuance rows, {orders} orders ({ws.max_row - 1} lines)")
    return path


def main() -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("path", type=Path, nargs="?", default=Path("data/uniforms.xlsx"))
    parser.add_argument("--count", type=int, default=48)
    args = parser.parse_args()
    build(args.path, count=args.count)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
